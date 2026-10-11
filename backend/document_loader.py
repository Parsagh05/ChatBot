"""Extract text from the company documents in knowledge_base/raw_docs."""

import logging
import re
from pathlib import Path
from typing import Any, Dict, List

from config import settings

logger = logging.getLogger(__name__)

SUPPORTED_EXTENSIONS = {".docx", ".pdf", ".txt", ".md"}


def load_raw_documents() -> List[Dict[str, Any]]:
    """One document per supported file in RAW_DOCS_DIR."""
    folder = Path(settings.RAW_DOCS_DIR)
    if not folder.is_dir():
        logger.warning("Raw documents folder not found: %s", folder)
        return []

    documents = []
    for path in sorted(folder.iterdir()):
        if path.suffix.lower() not in SUPPORTED_EXTENSIONS or path.name.startswith(("~$", ".")):
            continue
        try:
            text = extract_text(path)
        except Exception as exc:
            logger.error("Could not read %s: %s", path.name, exc)
            continue
        if not text.strip():
            logger.warning("No text extracted from %s", path.name)
            continue
        documents.append(
            {
                "id": f"doc:{path.name}",
                "source_type": "document",
                "title": document_title(path),
                "url": "",
                "categories": [],
                "text": text,
            }
        )
    return documents


def document_title(path: Path) -> str:
    """'01- درباره شرکت تابان انرژی- Rev02.docx' -> 'درباره شرکت تابان انرژی'."""
    title = re.sub(r"^[\d\s\-_.]+", "", path.stem)
    title = re.sub(r"[\s\-_]*rev\s*\d+\s*$", "", title, flags=re.IGNORECASE)
    return title.strip() or path.stem


def extract_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        return _docx_text(path)
    if suffix == ".pdf":
        return _pdf_text(path)
    return path.read_text(encoding="utf-8", errors="replace")


def _docx_text(path: Path) -> str:
    import docx

    document = docx.Document(str(path))
    lines = [p.text.strip() for p in document.paragraphs if p.text.strip()]
    for table in document.tables:
        for row in table.rows:
            cells = []
            for cell in row.cells:
                value = cell.text.strip()
                if value and value not in cells:  # merged cells repeat their text
                    cells.append(value)
            if cells:
                lines.append(" | ".join(cells))
    return "\n".join(lines)


def _pdf_text(path: Path) -> str:
    import pymupdf

    # default block order keeps right-to-left lines in reading order; sort=True reverses Persian words
    with pymupdf.open(str(path)) as pdf:
        pages = [page.get_text("text") for page in pdf]
    lines = [re.sub(r"[ \t]+", " ", line).strip() for page in pages for line in page.splitlines()]
    return "\n".join(line for line in lines if line)
