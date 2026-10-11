import json
from pathlib import Path

import pytest

from config import Settings, settings
from document_loader import document_title, load_raw_documents
from knowledge_store import build_chunks, build_link_resolver, collect_sources, split_text
from site_crawler import html_to_text


# ─── Chunking ────────────────────────────────────────────────────────────────


def test_split_text_respects_size_and_keeps_every_word():
    paragraphs = [" ".join(f"کلمه{p}_{w}" for w in range(40)) + "." for p in range(12)]
    text = "\n".join(paragraphs)
    chunks = split_text(text, chunk_size=300, chunk_overlap=60)
    assert len(chunks) > 1
    assert all(len(chunk) <= 300 + 60 for chunk in chunks)
    joined = " ".join(chunks)
    assert all(word in joined for word in text.split())


def test_split_text_overlaps_consecutive_chunks():
    text = "\n".join(f"جمله شماره {i} درباره پنل خورشیدی است." for i in range(60))
    chunks = split_text(text, chunk_size=200, chunk_overlap=50)
    for previous, current in zip(chunks, chunks[1:]):
        assert previous.split("\n")[-1] in current


def test_split_text_breaks_a_single_long_sentence():
    chunks = split_text("واژه " * 500, chunk_size=200, chunk_overlap=0)
    assert all(len(chunk) <= 200 for chunk in chunks)


def test_split_text_empty():
    assert split_text("", 800, 150) == []


def test_html_to_text_drops_scripts_and_forms():
    html = "<h2>عنوان</h2><p>متن <b>مهم</b></p><script>alert(1)</script><form><input value='x'>ارسال</form>"
    assert html_to_text(html) == "عنوان\nمتن\nمهم"


@pytest.mark.parametrize(
    "filename, title",
    [
        ("01- درباره شرکت تابان انرژی- Rev02.docx", "درباره شرکت تابان انرژی"),
        ("03- 07- فروشگاه.docx", "فروشگاه"),
        ("گزارش.pdf", "گزارش"),
    ],
)
def test_document_title(filename, title):
    assert document_title(Path(filename)) == title


# ─── Sources ─────────────────────────────────────────────────────────────────


@pytest.fixture
def knowledge_dir():
    base = Path(settings.KNOWLEDGE_BASE_DIR)
    raw = base / "raw_docs"
    raw.mkdir(parents=True, exist_ok=True)
    yield base
    for path in sorted(base.rglob("*"), reverse=True):
        path.unlink() if path.is_file() else path.rmdir()


def test_collect_sources_reads_all_kinds_and_drops_duplicates(knowledge_dir):
    import docx
    import pymupdf

    (knowledge_dir / "QA.json").write_text(
        json.dumps([{"question_id": "faq_01", "question": "آدرس کجاست؟", "answer": "ساری", "source_url": "tabenergy.ir/contact"}], ensure_ascii=False),
        encoding="utf-8",
    )
    (knowledge_dir / "knowledge_base.json").write_text(
        json.dumps(
            [
                {"chunk_id": "faq_01", "chunk_text": "سوال: آدرس کجاست؟\nپاسخ: ساری", "metadata": {"source_title": "آدرس کجاست؟"}},
                {"chunk_id": "hybrid_01", "chunk_text": "سیستم هیبرید ترکیبی از برق شبکه و باتری است.", "metadata": {"source_title": "هیبرید", "source_url": "https://tabenergy.ir/service/"}},
            ],
            ensure_ascii=False,
        ),
        encoding="utf-8",
    )
    document = docx.Document()
    document.add_paragraph("درباره شرکت")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "تلفن", "۰۱۱"
    document.save(knowledge_dir / "raw_docs" / "01- معرفی.docx")
    pdf = pymupdf.open()
    pdf.new_page().insert_text((72, 72), "Solar report 2025")
    pdf.save(knowledge_dir / "raw_docs" / "report.pdf")

    sources = collect_sources()
    by_type = {}
    for source in sources:
        by_type.setdefault(source["source_type"], []).append(source)

    assert [s["title"] for s in by_type["faq"]] == ["آدرس کجاست؟"]
    assert [s["title"] for s in by_type["curated"]] == ["هیبرید"]  # the copy of the FAQ entry is dropped
    docs = {s["title"]: s["text"] for s in by_type["document"]}
    assert docs["معرفی"] == "درباره شرکت\nتلفن | ۰۱۱"
    assert "Solar report 2025" in docs["report"]
    assert load_raw_documents()  # same loader the admin "Rebuild index" uses


def test_build_chunks_resolves_urls_and_keeps_faq_whole(knowledge_dir):
    long_answer = "پاسخ " * 400
    (knowledge_dir / "QA.json").write_text(
        json.dumps([{"question_id": "q", "question": "سوال؟", "answer": long_answer, "source_url": "tabenergy.ir/contact"}], ensure_ascii=False),
        encoding="utf-8",
    )
    sources = collect_sources()
    chunks = build_chunks(sources, build_link_resolver(sources))
    assert len(chunks) == 1
    assert chunks[0]["metadata"]["source_url"] == "https://tabenergy.ir/contact"
    assert chunks[0]["embed_text"].startswith("سوال؟\n")


# ─── Settings overrides ──────────────────────────────────────────────────────


def test_runtime_overrides_are_loaded_and_cast():
    path = settings.overrides_path
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps({"TOP_K": "9", "TEMPERATURE": 0.1, "NOT_A_SETTING": 1}), encoding="utf-8")
    fresh = Settings()
    assert fresh.TOP_K == 9 and fresh.TEMPERATURE == 0.1
    assert not hasattr(fresh, "NOT_A_SETTING")
